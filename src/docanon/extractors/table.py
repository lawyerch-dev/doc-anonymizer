"""表格抽取: XLSX / CSV。每个单元格一个 Block。"""
from __future__ import annotations

import csv
from pathlib import Path

from ..models import Block, ExtractedDoc
from .base import Extractor


class TableExtractor(Extractor):
    name = "table"
    extensions = (".xlsx", ".csv")

    def extract(self, path: Path) -> ExtractedDoc:
        suffix = path.suffix.lower()
        if suffix == ".csv":
            blocks = self._csv(path)
        else:
            blocks = self._xlsx(path)
        return ExtractedDoc(
            source_path=str(path), blocks=blocks, meta={"format": "table"}
        )

    @staticmethod
    def _csv(path: Path) -> list[Block]:
        blocks: list[Block] = []
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as fh:
            for r, row in enumerate(csv.reader(fh), start=1):
                for c, value in enumerate(row, start=1):
                    if value.strip():
                        blocks.append(
                            Block(
                                block_id=f"r{r}c{c}",
                                text=value,
                                kind="cell",
                                locator={"row": r, "col": c},
                            )
                        )
        return blocks

    @staticmethod
    def _xlsx(path: Path) -> list[Block]:
        from openpyxl import load_workbook

        wb = load_workbook(str(path), read_only=True, data_only=True)
        blocks: list[Block] = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    value = str(cell.value).strip()
                    if not value:
                        continue
                    blocks.append(
                        Block(
                            block_id=f"{ws.title}!{cell.coordinate}",
                            text=value,
                            kind="cell",
                            locator={"sheet": ws.title, "cell": cell.coordinate},
                        )
                    )
        wb.close()
        return blocks
