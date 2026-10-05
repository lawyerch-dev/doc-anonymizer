"""文字类文档抽取: TXT / MD / DOCX。"""
from __future__ import annotations

from pathlib import Path

from ..models import Block, ExtractedDoc
from .base import Extractor


class TextFileExtractor(Extractor):
    name = "text_file"
    extensions = (".txt", ".md", ".markdown", ".text")

    def extract(self, path: Path) -> ExtractedDoc:
        text = path.read_text(encoding="utf-8", errors="replace")
        # 按段落切块, 便于定位与增量处理
        blocks: list[Block] = []
        for idx, para in enumerate(text.split("\n")):
            if para.strip() == "":
                continue
            blocks.append(
                Block(
                    block_id=f"p{idx}",
                    text=para,
                    kind="paragraph",
                    locator={"line": idx},
                )
            )
        return ExtractedDoc(source_path=str(path), blocks=blocks, meta={"format": "text"})


class DocxExtractor(Extractor):
    name = "docx"
    extensions = (".docx",)

    def extract(self, path: Path) -> ExtractedDoc:
        from docx import Document  # python-docx

        doc = Document(str(path))
        blocks: list[Block] = []
        for idx, para in enumerate(doc.paragraphs):
            if not para.text.strip():
                continue
            blocks.append(
                Block(
                    block_id=f"p{idx}",
                    text=para.text,
                    kind="paragraph",
                    locator={"paragraph": idx},
                )
            )
        # 表格单元格
        for ti, table in enumerate(doc.tables):
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells):
                    for pi, para in enumerate(cell.paragraphs):
                        if not para.text.strip():
                            continue
                        blocks.append(
                            Block(
                                block_id=f"t{ti}r{ri}c{ci}p{pi}",
                                text=para.text,
                                kind="cell",
                                locator={"t": ti, "r": ri, "c": ci, "p": pi},
                            )
                        )
        return ExtractedDoc(source_path=str(path), blocks=blocks, meta={"format": "docx"})
