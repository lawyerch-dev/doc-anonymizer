"""把锚点还原回原文 —— **在原格式上就地回填**, 不重建文档。

为什么必须逐格式写, 不能"读成纯文本 → 替换 → 写回": docx 的样式挂在 run 的属性上
(`w:rPr`), xlsx 挂在单元格上。一旦读成纯文本, 这些属性就没了; 重新构造文档更是另一份文档了。
所以这里的做法是: **打开 AI 交回来的那个文件本身, 只改文字槽的内容, 一个属性都不碰**。
实测: 锚点被拆到 3 个 run 也能回填, 字体/字号/颜色/下划线完整保留。

定位一律靠**锚点文本**, 绝不靠位置(段落序号 / 单元格坐标): AI 会重排段落、合并单元格、
改写句子, 脱敏时记下的位置早就失效了。锚点天生对重排免疫 —— 这正是它存在的理由。
"""
from __future__ import annotations

import csv
import html
import re
from dataclasses import dataclass, field
from pathlib import Path

# 锚点形状: {{TYPE_N}}。TYPE 允许下划线, N 是数字。
# 匹配时**故意宽松**: AI 可能在锚点两侧加空格、改大小写、套上 markdown 强调号或反引号,
# 甚至把 {{ 转义成 HTML 实体。这些都是"同一个锚点被包了一层", 不该判成还原失败。
# 匹配出来后再用规范化形式(大写、去掉内部空白)去查表 —— 查不到就是未知锚点, 不猜。
_ANCHOR_RE = re.compile(r"\{\{\s*([A-Za-z_]+)\s*_(\d+)\s*\}\}")
# AI 有时把 {{ }} 转义成 HTML 实体({{ 本身不太会被转义, 但 &#123; 这类会出现)。
_ENTITY_RE = re.compile(r"&#?\w+;")


def normalize_anchor(entity_type: str, index: str) -> str:
    return f"{{{{{entity_type.strip().upper()}_{int(index)}}}}}"


def is_anchor(value: str) -> bool:
    """这个替换值是不是可寻址锚点(能被 restore 认出来)。

    pipeline 用它在账本里只记锚点 —— `mask`/`redact` 的替换值也进账本, 但它们还原不了,
    记进去只会让对账报告把一堆注定还原不回来的东西报成「缺失」。
    """
    text = value.strip()
    return bool(_ANCHOR_RE.fullmatch(text)) or bool(re.fullmatch(r"<[A-Za-z_]+_\d+>", text))


@dataclass
class RestoreResult:
    """一次还原的账: 换了多少、漏了什么、多出了什么。"""

    output: Path | None = None
    restored: dict[str, int] = field(default_factory=dict)      # 锚点 -> 还原次数
    unknown: list[str] = field(default_factory=list)            # 文中出现但映射表里没有
    missing: list[str] = field(default_factory=list)            # 该文件本该有、成品里没有
    scanned_units: int = 0

    @property
    def total_restored(self) -> int:
        return sum(v for k, v in self.restored.items() if not k.startswith("("))

    @property
    def complete(self) -> bool:
        return not self.unknown and not self.missing

    def to_dict(self) -> dict:
        return {
            "output": str(self.output) if self.output else None,
            "restored": {k: v for k, v in sorted(self.restored.items())},
            "total_restored": self.total_restored,
            "unknown": sorted(set(self.unknown)),
            "missing": sorted(set(self.missing)),
            "scanned_units": self.scanned_units,
            "complete": self.complete,
        }


class RestoreError(RuntimeError):
    """这份产物还原不了(格式不支持 / 读不了)。消息直接给用户看。"""


def _offset_in_original(original: str, probe: str, probe_pos: int) -> int:
    """unescape 之后 probe 的第 probe_pos 个字符, 在 original 里的下标。

    逐字符比而不是做比例换算: 实体替换让长度变化不均匀, 比例算出来的偏移会切错位置。
    只在真的发生过实体替换时才走到这里(常见路径 probe 与 original 同一个对象, 直接跳过)。
    """
    if probe_pos <= 0:
        return 0
    o = p = 0
    while o < len(original) and p < probe_pos:
        if original[o] == "&":
            m = _ENTITY_RE.match(original, o)
            if m:
                o = m.end()   # 整个实体在 probe 里算一个字符
                p += 1
                continue
        o += 1
        p += 1
    return o


def _collect(text: str, table: dict[str, str]) -> tuple[list[tuple[int, int, str]], list[str], dict[str, int]]:
    """在一段文字里找锚点。

    返回 (替换 span, 未知锚点, 锚点 → 命中次数)。命中不到的**不返回** —— 交给调用方记进
    unknown, 绝不在这里猜一个原文塞回去: AI 幻觉出的 {{PERSON_99}} 若被当成真锚点, 等于把某个
    无关的人名插进了不该出现的位置。
    """
    probe = html.unescape(text)
    plain = probe is text
    repls: list[tuple[int, int, str]] = []
    unknown: list[str] = []
    counts: dict[str, int] = {}
    for m in _ANCHOR_RE.finditer(probe):
        key = normalize_anchor(m.group(1), m.group(2))
        original = table.get(key)
        if original is None:
            unknown.append(key)
            continue
        start = m.start() if plain else _offset_in_original(text, probe, m.start())
        end = m.end() if plain else _offset_in_original(text, probe, m.end())
        repls.append((start, end, original))
        counts[key] = counts.get(key, 0) + 1
    return repls, unknown, counts


def _table_from(reverse: dict[str, str]) -> dict[str, str]:
    """反向表 → "规范锚点 → 原文"。

    只收能规范化的锚点(也就是 `anchor` / `placeholder` 策略产出的)。假名(`林芳`)与
    打码值(`138****5678`)不是锚点, 不走这条路径 —— 它们在纯文本 restore 里按字面替换,
    详见 `restore_literals`。
    """
    table: dict[str, str] = {}
    for repl, original in reverse.items():
        if not repl:
            continue  # remove 策略: 空串不是锚点, 拿去 replace 会把原文插到每个字符之间
        m = _ANCHOR_RE.fullmatch(repl.strip())
        if m:
            table[normalize_anchor(m.group(1), m.group(2))] = original
        else:
            m2 = re.fullmatch(r"<([A-Za-z_]+)_(\d+)>", repl.strip())
            if m2:  # 老版 placeholder 策略 `<TYPE_N>`, 历史 mapping.json 里躺着
                table[normalize_anchor(m2.group(1), m2.group(2))] = original
    return table


# ---------------------------------------------------------------- 纯文本 / csv

def _restore_text(src: Path, out: Path, table: dict[str, str], literals: list[tuple[str, str]]) -> RestoreResult:
    try:
        text = src.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as exc:
        raise RestoreError(f"读不了 {src.name}: 不是 UTF-8 文本") from exc
    repls, unknown, counts = _collect(text, table)
    out.write_text(_apply_literals(_apply(text, repls), literals), encoding="utf-8")
    return _finish(out, unknown, counts, units=1)


def _restore_csv(src: Path, out: Path, table: dict[str, str], literals: list[tuple[str, str]]) -> RestoreResult:
    with src.open("r", encoding="utf-8-sig", errors="replace", newline="") as fh:
        rows = list(csv.reader(fh))
    unknown: list[str] = []
    counts: dict[str, int] = {}
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            repls, unk, hits = _collect(value, table)
            unknown.extend(unk)
            if not repls and not unk:
                continue
            rows[r][c] = _apply_literals(_apply(value, repls), literals)
            for key, n in hits.items():
                counts[key] = counts.get(key, 0) + n
    with out.open("w", encoding="utf-8-sig", newline="") as fh:
        csv.writer(fh).writerows(rows)
    return _finish(out, unknown, counts, units=1)


# ---------------------------------------------------------------- docx

def _restore_docx(src: Path, out: Path, table: dict[str, str], _literals) -> RestoreResult:
    """docx 还原: 逐段落把锚点换回原文, **只改 run 的文字, run 属性一律不碰**。

    锚点被 Word/AI 拆到多个 run 里是常态(拼写检查、格式刷、往返转换都会拆), 所以不能
    "每个 run 各自替换" —— 那会一个都匹配不上。做法是: 先把整段的文字拼出来找锚点, 拿到
    全段坐标系里的 span, 再把 span 映射回各 run 的区间改写。这与 `writers._redact_paragraph`
    是同一套偏移算法的逆运算(那边是原文→锚点, 这边是锚点→原文)。
    """
    from docx import Document

    from ..docx_walk import iter_paragraphs, text_runs

    d = Document(str(src))
    unknown: list[str] = []
    counts: dict[str, int] = {}
    n_repl = 0
    n_para = 0
    for para in iter_paragraphs(d):
        runs = text_runs(para)
        if not runs:
            continue
        text = "".join(r.text for r in runs)
        repls, unk, hits = _collect(text, table)
        unknown.extend(unk)
        if not repls:
            continue
        for key, n in hits.items():
            counts[key] = counts.get(key, 0) + n
        n_para += 1
        n_repl += _rewrite_runs(runs, repls)
    if not n_para:
        # 一个锚点都没有: 别让 python-docx 重写一遍文件(会顺手丢掉它不认识的部件),
        # 原样复制字节。
        out.write_bytes(src.read_bytes())
    else:
        d.save(str(out))
    return _finish(out, unknown, counts, units=n_para)


def _rewrite_runs(runs, repls: list[tuple[int, int, str]]) -> int:
    """把 (start, end, 原文) 这些 span 写回 runs, 返回实际替换次数。

    替换值落在**起点所在的那个 run** 上, 后续被覆盖到的 run 只清掉重叠部分 —— 它们的字体
    属于被替换掉的那段文字, 留着会串味。起点 run 的属性(字体/字号/颜色)完整保留, 那是
    还原出来的原文该继承的样式。
    """
    ranges: list[tuple[int, int]] = []
    pos = 0
    for r in runs:
        ranges.append((pos, pos + len(r.text)))
        pos += len(r.text)
    ordered = sorted(repls)
    done = 0
    for i, (a, b) in enumerate(ranges):
        rt = runs[i].text
        out: list[str] = []
        cursor = 0
        for s, e, rep in ordered:
            if e <= a or s >= b:
                continue
            ls = max(s, a) - a
            le = min(e, b) - a
            if ls > cursor:
                out.append(rt[cursor:ls])
            if a <= s < b:  # 替换文本落在本 run 起点
                out.append(rep)
                done += 1
            cursor = max(cursor, le)
        out.append(rt[cursor:])
        new = "".join(out)
        if new != rt:
            runs[i].text = new
    return done


# ---------------------------------------------------------------- xlsx

def _restore_xlsx(src: Path, out: Path, table: dict[str, str], _literals) -> RestoreResult:
    """xlsx 还原: 逐单元格替换。单元格样式挂在 cell 上, 改 value 不影响它。"""
    from openpyxl import load_workbook

    wb = load_workbook(str(src))
    unknown: list[str] = []
    counts: dict[str, int] = {}
    n_cell = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                repls, unk, hits = _collect(cell.value, table)
                unknown.extend(unk)
                if not repls:
                    continue
                cell.value = _apply(cell.value, repls)
                n_cell += 1
                for key, n in hits.items():
                    counts[key] = counts.get(key, 0) + n
    if not n_cell:
        out.write_bytes(src.read_bytes())
    else:
        wb.save(str(out))
    return _finish(out, unknown, counts, units=n_cell)


# ---------------------------------------------------------------- 公共

def _apply(text: str, repls: list[tuple[int, int, str]]) -> str:
    for start, end, rep in sorted(repls, key=lambda x: x[0], reverse=True):
        text = text[:start] + rep + text[end:]
    return text


def _apply_literals(text: str, literals: list[tuple[str, str]]) -> str:
    for repl, original in literals:
        if repl and repl in text:
            text = text.replace(repl, original)
    return text


def _finish(out: Path, unknown: list[str], counts: dict[str, int], units: int) -> RestoreResult:
    return RestoreResult(output=out, restored=counts, unknown=unknown, scanned_units=units)


_ROUTES = {
    ".docx": _restore_docx,
    ".xlsx": _restore_xlsx,
    ".csv": _restore_csv,
    ".txt": _restore_text,
    ".md": _restore_text,
    ".markdown": _restore_text,
}

# 这些格式**原理上还原不了**: 脱敏时命中页整页栅格化, 原文字层已被物理删除。
# 早于该版本生成的产物若用的是"盖黑块但留文字层", 也不在支持范围内 —— 宁可明确拒绝,
# 也不要交出一份看起来还原了、其实原文还留在里面的文件。
_UNRESTORABLE = {
    ".pdf": "PDF 命中页在脱敏时已整页栅格化, 原文字层被删除, 拿不回原文。"
            "要走「脱敏 → AI → 还原」请先用 docx。",
    ".png": "图片上的敏感信息是涂黑, 像素里没有原文, 还原不了。",
    ".jpg": "同上。",
    ".jpeg": "同上。",
    ".bmp": "同上。",
    ".tif": "同上。",
    ".tiff": "同上。",
    ".webp": "同上。",
}


def restore_file(
    src: Path,
    out: Path,
    reverse: dict[str, str],
    expected: dict[str, int] | None = None,
) -> RestoreResult:
    """按扩展名路由还原。

    `reverse`   —— mapping.json 的反向表(替换值 → 原文)。
    `expected`  —— **这一个文件**脱敏时写进去的锚点及次数(来自 manifest 的 `anchors` 字段)。
                   有了它才能判定"少了锚点"; 没有就只报 unknown, 不猜谁是漏的 ——
                   整张映射表覆盖所有源文件, 拿全表来对账会把别的文件的锚点误报成缺失。
    """
    suffix = src.suffix.lower()
    if suffix in _UNRESTORABLE:
        raise RestoreError(_UNRESTORABLE[suffix])
    route = _ROUTES.get(suffix)
    if route is None:
        raise RestoreError(
            f"不认识的格式 {suffix or '(无扩展名)'}。可还原: "
            f"{' '.join(sorted(_ROUTES))}; 不可还原: {' '.join(sorted(_UNRESTORABLE))}。"
        )
    table = _table_from(reverse)
    literals = _literals_from(reverse)
    out.parent.mkdir(parents=True, exist_ok=True)
    res = route(src, out, table, literals)
    if expected:
        res.missing = [k for k in expected if k not in res.restored]
    return res


def _literals_from(reverse: dict[str, str]) -> list[tuple[str, str]]:
    """非锚点形态的替换值(假名、打码值), 按字面做替换。

    这些**只用于纯文本产物**: 它们不是可寻址锚点, 在 docx 里可能跨 run, 按字面全局替换
    反而会打到不相干的地方。docx/xlsx 路线传空表, 只认锚点。「交 AI 再还原」这条链用
    `anchor` 策略, 天然全是锚点, 不受影响。

    另有一道安全闸: 纯符号类替换值(`redact` 的 `**`)必须排除 —— markdown 的加粗也是 `**`,
    拿它去全局替换等于把整篇文档的加粗标记换成某个当事人姓名。判据是"含至少一个字母/数字/汉字"。
    """
    out: list[tuple[str, str]] = []
    for repl, original in sorted(reverse.items(), key=lambda kv: -len(kv[0])):
        if not repl or _ANCHOR_RE.fullmatch(repl.strip()) or re.fullmatch(r"<[A-Za-z_]+_\d+>", repl.strip()):
            continue
        if not any(ch.isalnum() for ch in repl):
            continue
        out.append((repl, original))
    return out