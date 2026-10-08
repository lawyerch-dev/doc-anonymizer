"""原位回写: 把脱敏结果写回**原格式**(docx/xlsx/csv/pdf/图片), 供 before/after 对比预览。"""
from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path

from docanon_contract import Block, ExtractedDoc


@dataclass
class BlockRedaction:
    block: Block
    replacements: list[tuple[int, int, str]] = field(default_factory=list)
    new_text: str = ""

    @property
    def changed(self) -> bool:
        return bool(self.replacements)


def write_output(doc: ExtractedDoc, reds: list[BlockRedaction], out_path: Path) -> list[str]:
    fmt = doc.meta.get("format")
    if fmt == "docx":
        return _docx(doc, reds, out_path)
    if fmt == "table":
        return _table(doc, reds, out_path)
    if fmt == "pdf":
        return _pdf(doc, reds, out_path)
    if fmt == "image":
        return _image(doc, reds, out_path)
    return _text(doc, reds, out_path)


# ---------------- 纯文本 / Markdown ----------------
def _text(doc: ExtractedDoc, reds: list[BlockRedaction], out_path: Path) -> list[str]:
    src = Path(doc.source_path)
    lines = src.read_text(encoding="utf-8", errors="replace").split("\n")
    for r in reds:
        if r.changed:
            i = r.block.locator.get("line")
            if i is not None and 0 <= i < len(lines):
                lines[i] = r.new_text
    out_path.write_text("\n".join(lines), encoding="utf-8")
    return [str(out_path)]


# ---------------- DOCX ----------------
def _redact_paragraph(paragraph, repls: list[tuple[int, int, str]]) -> None:
    """按字符偏移改写段落文字。

    `repls` 的偏移来自 `docx_walk.paragraph_text` 拼出来的文字, 所以这里的 run 序列也必须
    来自 `docx_walk.text_runs` —— 换成 `paragraph.runs` 就会漏掉超链接/修订插入里的 run,
    偏移整体错位(实测: 目标原样留着, 替换值插到别的字上)。
    """
    from ..docx_walk import text_runs

    runs = text_runs(paragraph)
    if not runs:
        return
    ranges: list[tuple[int, int]] = []
    pos = 0
    for run in runs:
        ranges.append((pos, pos + len(run.text)))
        pos += len(run.text)
    ordered = sorted(repls)
    for i, (a, b) in enumerate(ranges):
        rt = runs[i].text
        out: list[str] = []
        cursor = 0
        for s, e, rep in ordered:
            if e <= a or s >= b:
                continue
            ls = max(s, a) - a
            le = min(e, b) - a
            if ls > cursor:
                out.append(rt[cursor:ls])
            if a <= s < b:  # 替换文本落在本 run 起点
                out.append(rep)
            cursor = max(cursor, le)
        out.append(rt[cursor:])
        new = "".join(out)
        if new != rt:
            runs[i].text = new


def _scrub_link_plumbing(paragraph, pairs: list[tuple[str, str]]) -> None:
    """把本段换掉的原文也从**链接目标**里抹掉。

    超链接的可见文字改了、`mailto:`/URL 里还留着原文, 这份文件里就仍然搜得到它 —— 对
    一个"不许漏"的工具来说等于没脱。指令文字(`w:instrText`, 域代码写法)与关系目标
    (`word/_rels/document.xml.rels`) 都是原文会藏的地方。

    只替换本段实际换过的原文, 不做通用 URL 扫描(那属于检测器该干的事, 还没做)。
    """
    from docx.oxml.ns import qn

    element = paragraph._p
    for instr in element.iter(qn("w:instrText")):
        text = instr.text or ""
        new = text
        for original, repl in pairs:
            new = new.replace(original, repl)
        if new != text:
            instr.text = new

    rels = paragraph.part.rels
    for link in element.iter(qn("w:hyperlink")):
        rid = link.get(qn("r:id"))
        rel = rels.get(rid) if rid else None
        if rel is None or not rel.is_external:
            continue
        target = rel.target_ref
        new = target
        for original, repl in pairs:
            new = new.replace(original, repl)
        if new != target:
            rel._target = new


def _docx(doc: ExtractedDoc, reds: list[BlockRedaction], out_path: Path) -> list[str]:
    from docx import Document

    from ..docx_walk import iter_paragraphs

    d = Document(doc.source_path)
    paragraphs = list(iter_paragraphs(d))
    for r in reds:
        if not r.changed:
            continue
        idx = r.block.locator.get("paragraph")
        if idx is None or not 0 <= idx < len(paragraphs):
            # 抽取与回写走同一个遍历, 对不上就是本程序的 bug: 宁可报错也别交出没脱敏的产物。
            raise RuntimeError(
                f"docx 段落定位失败(locator={r.block.locator!r}, 共 {len(paragraphs)} 段)"
            )
        para = paragraphs[idx]
        _redact_paragraph(para, r.replacements)
        _scrub_link_plumbing(
            para, [(r.block.text[s:e], rep) for s, e, rep in r.replacements if r.block.text[s:e]]
        )
    d.save(str(out_path))
    return [str(out_path)]


# ---------------- 表格 ----------------
def _table(doc: ExtractedDoc, reds: list[BlockRedaction], out_path: Path) -> list[str]:
    by_cell: dict[tuple, str] = {}
    for r in reds:
        if not r.changed:
            continue
        loc = r.block.locator
        if "row" in loc:
            by_cell[("csv", loc["row"], loc["col"])] = r.new_text
        elif "cell" in loc:
            by_cell[(loc.get("sheet"), loc["cell"])] = r.new_text

    if out_path.suffix.lower() == ".csv":
        src = Path(doc.source_path)
        with src.open("r", encoding="utf-8-sig", errors="replace", newline="") as fh:
            rows = list(csv.reader(fh))
        for (kind, r, c), text in by_cell.items():
            if kind == "csv" and 1 <= r <= len(rows) and 1 <= c <= len(rows[r - 1]):
                rows[r - 1][c - 1] = text
        with out_path.open("w", encoding="utf-8-sig", newline="") as fh:
            csv.writer(fh).writerows(rows)
        return [str(out_path)]

    from openpyxl import load_workbook

    wb = load_workbook(doc.source_path)
    for (sheet, coord), text in by_cell.items():
        ws = wb[sheet] if sheet in wb.sheetnames else wb.active
        ws[coord] = text
    wb.save(str(out_path))
    return [str(out_path)]


# ---------------- PDF(文字层 + 扫描页) ----------------
_SCALE = 2.0


def _text_page_boxes(pdf, pno: int, repls, scale: float) -> list[tuple]:
    page = pdf[pno]
    width_pt, height_pt = page.get_size()
    textpage = page.get_textpage()
    boxes: list[tuple] = []
    for s, e, _rep in repls:
        char_boxes = []
        for ci in range(s, e):
            try:
                left, bottom, right, top = textpage.get_charbox(ci)
            except Exception:
                continue
            char_boxes.append((left, bottom, right, top))
        if char_boxes:
            x0 = min(b[0] for b in char_boxes) * scale
            x1 = max(b[2] for b in char_boxes) * scale
            y0 = (height_pt - max(b[3] for b in char_boxes)) * scale
            y1 = (height_pt - min(b[1] for b in char_boxes)) * scale
            boxes.append((x0, y0, x1, y1))
    return boxes


def _char_width(ch: str) -> float:
    # CJK/全角约 2 个宽度单位, ASCII 约 1 个
    return 2.0 if ord(ch) > 0x2E7F else 1.0


def _frac(text: str, idx: int) -> float:
    total = sum(_char_width(c) for c in text) or 1.0
    pre = sum(_char_width(c) for c in text[:idx])
    return max(0.0, min(1.0, pre / total))


def _sub_boxes(bbox, text: str, repls) -> list[tuple]:
    """OCR 只给整行 bbox; 按**字符宽度比例**(CJK=2, ASCII=1)估子矩形, 只涂敏感片段。"""
    x0, y0, x1, y1 = bbox
    width = x1 - x0
    boxes = []
    for s, e, _rep in repls:
        f0 = _frac(text, s)
        f1 = _frac(text, e)
        if f1 <= f0:
            f1 = min(1.0, f0 + 0.02)
        boxes.append((x0 + width * f0, y0, x0 + width * f1, y1))
    return boxes


def _pdf(doc: ExtractedDoc, reds: list[BlockRedaction], out_path: Path) -> list[str]:
    """命中的页栅格化后涂黑, **没命中的页原样保留**(不重渲染)。

    为什么不给命中页做"保留文字层 + 盖黑块": 那样敏感文字仍可被复制/搜索, 等于没脱敏 —— 宁可
    让这一页变成位图。代价是这一页不可再编辑、体积变大; 但没命中的页没必要陪着一起变位图,
    所以它们直接从原文件导入(矢量文字、体积都不变)。
    """
    import pypdfium2 as pdfium
    from PIL import ImageDraw

    scanned = {p["page"]: p["image"] for p in doc.meta.get("_scanned_images", [])}
    # 收集每页要涂黑的框
    text_repls: dict[int, list] = {}
    img_boxes: dict[int, list] = {}
    for r in reds:
        if not r.changed:
            continue
        loc = r.block.locator
        pno = loc.get("page", 0)
        if "bbox" in loc:
            img_boxes.setdefault(pno, []).extend(
                _sub_boxes(loc["bbox"], r.block.text, r.replacements)
            )
        else:
            text_repls.setdefault(pno, []).extend(r.replacements)

    pdf = pdfium.PdfDocument(doc.source_path)
    out = pdfium.PdfDocument.new()
    for pno in range(len(pdf)):
        boxes = img_boxes.get(pno, [])
        repls = text_repls.get(pno, [])
        if not boxes and not repls:
            out.import_pages(pdf, [pno])  # 没命中: 原样搬过来
            continue
        img = scanned[pno] if pno in scanned else pdf[pno].render(scale=_SCALE).to_pil()
        draw = ImageDraw.Draw(img)
        for box in boxes:
            draw.rectangle(box, fill="black")
        for box in _text_page_boxes(pdf, pno, repls, _SCALE):
            draw.rectangle(box, fill="black")
        # 位图按原文页面尺寸铺满: 矩阵给的是页面单位(pt), 不是 1/scale —— 给错等于没涂黑
        width_pt, height_pt = pdf[pno].get_size()
        bitmap = pdfium.PdfImage.new(out)
        bitmap.set_bitmap(pdfium.PdfBitmap.from_pil(img.convert("RGB")))
        bitmap.set_matrix(pdfium.PdfMatrix().scale(width_pt, height_pt))
        page = out.new_page(width_pt, height_pt)
        page.insert_obj(bitmap)
        page.gen_content()

    if len(out) == 0:
        return []
    out.save(str(out_path))
    return [str(out_path)]


# ---------------- 图片 ----------------
def _image(doc: ExtractedDoc, reds: list[BlockRedaction], out_path: Path) -> list[str]:
    from PIL import Image, ImageDraw

    img = Image.open(doc.image_path).convert("RGB")
    draw = ImageDraw.Draw(img)
    for r in reds:
        if r.changed and "bbox" in r.block.locator:
            for box in _sub_boxes(r.block.locator["bbox"], r.block.text, r.replacements):
                draw.rectangle(box, fill="black")
    img.save(out_path)
    return [str(out_path)]
