"""DOCX 回写: 断言产物里的**文字**, 不是只数命中数。

踩过的坑: `paragraph.text` 含超链接里的文字, 而 `paragraph.runs` 不含 —— 拿后者套前者的
字符偏移, 结果是超链接里的邮箱原样留着、旁边的字被换成替换值, 而 manifest 报 `ok`。
所以这里每条都同时断言"原文不在了"和"旁边的字没被动过"。
"""
from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from docanon_core.config import load_config
from docanon_core.pipeline import process_file
from docanon_core.redaction.mapping import MappingStore

PHONE = "13800001111"
EMAIL = "a@b.com"
ID = "11010119900307721X"


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _redact(docx: Path, tmp_path: Path):
    cfg = load_config()
    res = process_file(docx, tmp_path / "out", cfg, MappingStore())
    return res, Path(res.output_path)


def _doc_text(path: Path) -> str:
    """产物里所有可见文字(直接读 XML, 免得又被 python-docx 的某个视图骗到)。"""
    xml = zipfile.ZipFile(path).read("word/document.xml").decode()
    return "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml))


def _rels_text(path: Path) -> str:
    return zipfile.ZipFile(path).read("word/_rels/document.xml.rels").decode()


def _paragraphs(path: Path) -> list[str]:
    return [p.text for p in Document(str(path)).paragraphs]


def test_pii_after_hyperlink_keeps_offsets(tmp_path):
    """超链接在前: 曾经会把替换值插到超链接旁、电话号码改烂。"""
    d = Document()
    p = d.add_paragraph("B 行: 邮箱 ")
    _add_hyperlink(p, EMAIL, f"mailto:{EMAIL}")
    p.add_run(f" 联系电话 {PHONE}")
    src = tmp_path / "a.docx"
    d.save(src)

    _, out = _redact(src, tmp_path)
    assert _paragraphs(out) == ["B 行: 邮箱 a***@b.com 联系电话 138****1111"]
    assert EMAIL not in _doc_text(out) and PHONE not in _doc_text(out)


def test_pii_inside_hyperlink_is_redacted_and_target_scrubbed(tmp_path):
    """敏感值本身就是超链接(Word 里点得开的邮箱/网址默认这样): 文字与链接目标都要干净。"""
    d = Document()
    p = d.add_paragraph("C 行: 邮箱 ")
    _add_hyperlink(p, EMAIL, f"mailto:{EMAIL}")
    p.add_run(" 已确认")
    src = tmp_path / "a.docx"
    d.save(src)

    _, out = _redact(src, tmp_path)
    assert _paragraphs(out) == ["C 行: 邮箱 a***@b.com 已确认"]
    # 可见文字换了、`mailto:` 里还留着原文, 这份文件里照样能搜到它
    assert EMAIL not in _doc_text(out)
    assert EMAIL not in _rels_text(out)


def test_pii_before_hyperlink_is_redacted(tmp_path):
    d = Document()
    p = d.add_paragraph(f"D 行: 电话 {PHONE} 邮箱 ")
    _add_hyperlink(p, "x@y.com", "mailto:x@y.com")
    p.add_run(" 完")
    src = tmp_path / "a.docx"
    d.save(src)

    _, out = _redact(src, tmp_path)
    assert _paragraphs(out) == ["D 行: 电话 138****1111 邮箱 x***@y.com 完"]
    assert PHONE not in _doc_text(out) and "x@y.com" not in _doc_text(out)


def test_span_split_across_runs_is_redacted(tmp_path):
    """Word 编辑会把一个号码拆进多个 run —— 拼接顺序错了同样会改错字。"""
    d = Document()
    p = d.add_paragraph("A 行: ")
    for piece in ("138", "0000", "1111"):
        p.add_run(piece)
    p.add_run(" 完")
    src = tmp_path / "a.docx"
    d.save(src)

    _, out = _redact(src, tmp_path)
    assert _paragraphs(out) == ["A 行: 138****1111 完"]


def test_revision_insert_runs_are_redacted(tmp_path):
    """修订插入(w:ins)里的文字也算正文 —— 当事人信息常是在修订态里补进去的。"""
    d = Document()
    p = d.add_paragraph("A 行: ")
    ins = OxmlElement("w:ins")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = f"电话 {PHONE}"
    run.append(t)
    ins.append(run)
    p._p.append(ins)
    src = tmp_path / "a.docx"
    d.save(src)

    _, out = _redact(src, tmp_path)
    assert PHONE not in _doc_text(out)
    assert "138****1111" in _doc_text(out)


def test_table_and_nested_table_cells_are_redacted(tmp_path):
    d = Document()
    outer = d.add_table(rows=1, cols=1)
    outer.rows[0].cells[0].text = f"外层单元格 {EMAIL}"
    inner = outer.rows[0].cells[0].add_table(rows=1, cols=1)
    inner.rows[0].cells[0].text = f"内层单元格 {ID}"
    d.add_paragraph("正文段落")
    src = tmp_path / "a.docx"
    d.save(src)

    _, out = _redact(src, tmp_path)
    text = _doc_text(out)
    assert EMAIL not in text and ID not in text
    assert "正文段落" in text


def test_content_control_paragraph_is_redacted(tmp_path):
    """内容控件(w:sdt)是 Word 模板/表单的常见写法。"""
    d = Document()
    d.add_paragraph("正文无敏感信息")
    sdt = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    p = OxmlElement("w:p")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = f"控件内: 联系电话 {PHONE}"
    run.append(t)
    p.append(run)
    content.append(p)
    sdt.append(content)
    d.element.body.append(sdt)
    src = tmp_path / "a.docx"
    d.save(src)

    _, out = _redact(src, tmp_path)
    text = _doc_text(out)
    assert PHONE not in text
    assert "控件内: 联系电话 138****1111" in text


def test_nothing_from_the_source_text_survives(tmp_path):
    """兜底: 上面所有排布塞进一个文件, 产物里一个原文都不许剩。"""
    d = Document()
    d.add_paragraph(f"顶层 {PHONE}")
    p = d.add_paragraph("链接 ")
    _add_hyperlink(p, EMAIL, f"mailto:{EMAIL}")
    d.add_table(rows=1, cols=1).rows[0].cells[0].text = f"表格 {ID}"
    src = tmp_path / "a.docx"
    d.save(src)

    res, out = _redact(src, tmp_path)
    blob = _doc_text(out) + _rels_text(out)
    for original in (PHONE, EMAIL, ID):
        assert original not in blob, f"{original} 还留在产物里"
    assert set(res.entity_counts) >= {"PHONE", "EMAIL", "ID_CARD"}
