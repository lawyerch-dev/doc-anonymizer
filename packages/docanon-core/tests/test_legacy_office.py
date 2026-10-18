"""旧格式端到端: 转换 -> 抽取 -> 脱敏 -> 回写为现代格式。"""
from __future__ import annotations

import tempfile
from pathlib import Path

import openpyxl
import pytest
from docx import Document

from docanon_core.config import load_config
from docanon_core.pipeline import process_file
from docanon_core.redaction.mapping import MappingStore
from _soffice_stub import install_fake_soffice


def _track_tmpdirs(monkeypatch, base: Path) -> list[Path]:
    """把 tempfile.mkdtemp 重定向到 tmp_path, 记录每个被建出来的临时目录。"""
    real = tempfile.mkdtemp
    made: list[Path] = []

    def fake(**kw):
        d = real(prefix=kw.get("prefix", "tmp"), dir=str(base))
        made.append(Path(d))
        return d

    monkeypatch.setattr(tempfile, "mkdtemp", fake)
    return made


def test_doc_becomes_redacted_docx(tmp_path, monkeypatch):
    install_fake_soffice(tmp_path, monkeypatch)
    src = tmp_path / "合同.doc"
    src.write_bytes(b"\xd0\xcf\x11\xe0 legacy")
    out = tmp_path / "out"
    out.mkdir()

    res = process_file(src, out, load_config(), MappingStore())

    assert res.converted_from == ".doc"
    assert res.output_format == "docx"
    produced = Path(res.output_path)
    assert produced.name == "【脱敏版】合同.doc.docx"
    text = "\n".join(p.text for p in Document(str(produced)).paragraphs)
    assert "13812340000" not in text, "转换后的 docx 没被脱敏"


def test_xls_keeps_xlsx_extension(tmp_path, monkeypatch):
    install_fake_soffice(tmp_path, monkeypatch)
    src = tmp_path / "表.xls"
    src.write_bytes(b"\xd0\xcf\x11\xe0 legacy")
    out = tmp_path / "out"
    out.mkdir()

    res = process_file(src, out, load_config(), MappingStore())

    assert res.output_format == "xlsx"
    assert Path(res.output_path).name == "【脱敏版】表.xls.xlsx"
    assert "13912345678" not in openpyxl.load_workbook(res.output_path).active["A1"].value


def test_temp_workdir_is_cleaned_on_success(tmp_path, monkeypatch):
    install_fake_soffice(tmp_path, monkeypatch)
    made = _track_tmpdirs(monkeypatch, tmp_path)
    src = tmp_path / "合同.doc"
    src.write_bytes(b"\xd0\xcf\x11\xe0 legacy")
    out = tmp_path / "out"
    out.mkdir()

    process_file(src, out, load_config(), MappingStore())

    assert made and not made[-1].exists(), "转换临时目录必须被清理"


def test_temp_workdir_is_cleaned_on_conversion_failure(tmp_path, monkeypatch):
    """转换失败也要收临时目录 —— 否则每个坏文件都漏一个目录。"""
    install_fake_soffice(tmp_path, monkeypatch, fail=True)
    made = _track_tmpdirs(monkeypatch, tmp_path)
    src = tmp_path / "合同.doc"
    src.write_bytes(b"\xd0\xcf\x11\xe0 legacy")
    out = tmp_path / "out"
    out.mkdir()

    with pytest.raises(RuntimeError):
        process_file(src, out, load_config(), MappingStore())

    assert made and not made[-1].exists(), "转换失败时临时目录泄漏了"
