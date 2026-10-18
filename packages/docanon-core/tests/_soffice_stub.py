"""测试用的假 soffice: 把预置的 docx/xlsx 拷成目标名, 让旧格式转换离线可测。

**只此一份**: 这段 shell 片段和 fixture 造法一旦在几个测试里各写一遍, 漂移时不会有任何测试报警
—— 它们会各自"自洽"地骗过去。四个调用方(转换单测、旧格式端到端、服务端、产物布局)都从这里取,
差异靠参数表达, 不靠复制。

参数:

- `real_documents`: 产物是真 docx/xlsx(带可脱敏的号码, 端到端要用)还是只认前缀的假字节;
- `fail`: soffice 直接 exit 1(验转换失败路径);
- `record_args`: 把 argv 追加到这个文件(验命令行约定, 如并发转换的 profile 隔离);
- `output_mode`: `source` = 产物按源名; `other` = 故意另起名 `out.<后缀>`; `both` = 两种都产出。
  `other` 验"别把别名/杂散文件当产物", `both` 验"多个候选时优先源名一致的那个"。

三种模式都会先落一个 `stray.tmp`: 它永远不该被当成产物, 于是这条过滤每次都被真的检验到。
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from docx import Document

_MODES = ("source", "other", "both")


def _shell(*, record: bool, mode: str) -> str:
    lines = [
        "#!/usr/bin/env bash",
        "set -euo pipefail",
        'echo "$@" >> "$ARGS_FILE"',
    ] if record else [
        "#!/usr/bin/env bash",
        "set -euo pipefail",
    ]
    lines += [
        'outdir=""; fmt=""',
        "while [ $# -gt 0 ]; do case \"$1\" in",
        '  --convert-to) fmt="$2"; shift 2;;',
        '  --outdir) outdir="$2"; shift 2;;',
        "  --headless|--norestore) shift;;",
        # 兜底分支也吃 `-env:UserInstallation=...`: 最后一个参数(真正的源文件)会覆盖它
        '  *) src="$1"; shift;;',
        "esac; done",
        # 杂散文件: 后缀不匹配, 永远不该被挑成产物
        ': > "$outdir/stray.tmp"',
        'stem="$(basename "$src")"; stem="${stem%.*}"',
    ]
    if mode in ("source", "both"):
        lines.append('cp "$FIXTURE_DIR/$fmt" "$outdir/$stem.$fmt"')
    if mode in ("other", "both"):
        lines.append('cp "$FIXTURE_DIR/$fmt" "$outdir/out.$fmt"')
    return "\n".join(lines) + "\n"


def install_fake_soffice(
    tmp_path: Path,
    monkeypatch,
    *,
    real_documents: bool = True,
    fail: bool = False,
    record_args: Path | None = None,
    output_mode: str = "source",
) -> Path:
    """装好假 soffice 与 fixture, 返回 fixture 目录。"""
    if output_mode not in _MODES:
        raise ValueError(f"未知的 output_mode: {output_mode!r}(可选: {', '.join(_MODES)})")

    fixture = tmp_path / "fixtures"
    fixture.mkdir()
    if real_documents:
        doc = Document()
        doc.add_paragraph("甲方 张三 手机 13812340000")
        doc.save(str(fixture / "docx"))
        wb = openpyxl.Workbook()
        wb.active["A1"] = "张三 13912345678"
        wb.save(str(fixture / "xlsx"))
    else:
        (fixture / "docx").write_bytes(b"PK\x03\x04 fake-docx")
        (fixture / "xlsx").write_bytes(b"PK\x03\x04 fake-xlsx")

    stub = tmp_path / "soffice"
    stub.write_text(
        "#!/usr/bin/env bash\nexit 1\n" if fail else _shell(record=record_args is not None, mode=output_mode),
        encoding="utf-8",
    )
    stub.chmod(0o755)
    monkeypatch.setenv("DOCANON_SOFFICE", str(stub))
    monkeypatch.setenv("FIXTURE_DIR", str(fixture))
    if record_args is not None:
        monkeypatch.setenv("ARGS_FILE", str(record_args))
    return fixture
