"""生成各常见格式的示例文档(含虚构敏感信息, 仅用于测试)。

用法: .venv/bin/python scripts/make_samples.py
"""
from __future__ import annotations

from pathlib import Path

SAMPLES = Path(__file__).resolve().parents[1] / "samples"
SAMPLES.mkdir(exist_ok=True)

TEXT = """关于「内部项目代号X」第二季度经营分析报告

汇报人：张三，联系电话 13812340000，邮箱 zhangsan@example.com。
身份证号 110101199003071234，银行卡 6222020200112233445。
公司地址：北京市朝阳区建国路88号，服务器 IP 192.168.1.100。
客户「某客户名」合同金额 860000 元，对接人李四，电话 13998765432。
本报告涉及内部项目代号X，请勿外传。
access_token = sk-abcdef1234567890ABCDEF
"""

CJK_FONT = "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"


def make_md() -> None:
    (SAMPLES / "sample.md").write_text("# 分析报告\n\n" + TEXT, encoding="utf-8")


def make_docx() -> None:
    from docx import Document

    doc = Document()
    for line in TEXT.splitlines():
        doc.add_paragraph(line)
    doc.save(str(SAMPLES / "sample.docx"))


def make_csv() -> None:
    rows = [
        ["姓名", "电话", "身份证", "备注"],
        ["张三", "13812340000", "110101199003071234", "内部项目代号X"],
        ["李四", "13998765432", "110101198802023456", "某客户名"],
    ]
    import csv

    with (SAMPLES / "sample.csv").open("w", encoding="utf-8-sig", newline="") as fh:
        csv.writer(fh).writerows(rows)


def make_xlsx() -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "客户"
    ws.append(["姓名", "电话", "身份证", "邮箱"])
    ws.append(["张三", "13812340000", "110101199003071234", "zhangsan@example.com"])
    ws.append(["李四", "13998765432", "110101198802023456", "lisi@example.com"])
    wb.save(str(SAMPLES / "sample.xlsx"))


def _cjk_font_name() -> str:
    """注册并返回一个**内嵌**的中文 TTF 字体名(避免依赖查看器的 cMap)。"""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    pdfmetrics.registerFont(TTFont("CJKEmbed", CJK_FONT))
    return "CJKEmbed"


def make_text_pdf() -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    font = _cjk_font_name()
    c = canvas.Canvas(str(SAMPLES / "sample_text.pdf"), pagesize=A4)
    c.setFont(font, 12)
    y = 800
    for line in TEXT.splitlines():
        c.drawString(50, y, line)
        y -= 22
    c.showPage()
    c.save()


def _render_image(text: str, size=(1000, 620)):
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.truetype(CJK_FONT, 30)
    y = 40
    for line in text.splitlines():
        draw.text((40, y), line, fill="black", font=font)
        y += 48
    return img


def make_scan_png() -> None:
    _render_image(TEXT).save(str(SAMPLES / "sample_scan.png"))


def make_scanned_pdf() -> None:
    img = _render_image(TEXT)
    img.convert("RGB").save(str(SAMPLES / "sample_scanned.pdf"), "PDF", resolution=150)


# ---- 多页样例 ----
PAGES = [
    "第 1 页 / 共 3 页\n汇报人：张三，联系电话 13812340000，邮箱 zhangsan@example.com。\n身份证号 110101199003071234。",
    "第 2 页 / 共 3 页\n客户「某客户名」合同金额 860000 元，对接人李四，电话 13998765432。\n银行卡 6222020200112233445。",
    "第 3 页 / 共 3 页\n服务器 IP 192.168.1.100，内部项目代号X 请勿外传。\naccess_token = sk-abcdef1234567890ABCDEF",
]


def make_multipage_text_pdf() -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    font = _cjk_font_name()
    c = canvas.Canvas(str(SAMPLES / "sample_multipage_text.pdf"), pagesize=A4)
    c.setFont(font, 12)
    for page in PAGES:
        y = 800
        for line in page.splitlines():
            c.drawString(50, y, line)
            y -= 22
        c.showPage()
    c.save()


def make_multipage_scanned_pdf() -> None:
    imgs = [_render_image(p).convert("RGB") for p in PAGES]
    imgs[0].save(str(SAMPLES / "sample_multipage_scanned.pdf"), save_all=True,
                 append_images=imgs[1:], resolution=150)


def make_multipage_docx() -> None:
    from docx import Document

    doc = Document()
    for i, page in enumerate(PAGES):
        for line in page.splitlines():
            doc.add_paragraph(line)
        if i < len(PAGES) - 1:
            doc.add_page_break()
    doc.save(str(SAMPLES / "sample_multipage.docx"))


def main() -> None:
    make_md()
    make_docx()
    make_csv()
    make_xlsx()
    make_text_pdf()
    make_scan_png()
    make_scanned_pdf()
    make_multipage_text_pdf()
    make_multipage_scanned_pdf()
    make_multipage_docx()
    for p in sorted(SAMPLES.glob("sample*")):
        print(f"生成 {p.name}  ({p.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
